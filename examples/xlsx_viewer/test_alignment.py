"""Headless tests for conservative alignment and Save As verification."""
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

from alignment import analyze, apply, fingerprint, ViewerError, validate_destination


class AlignmentTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.source = Path(self.directory.name) / 'source.xlsx'
        self.output = self.source.with_name('aligned.xlsx')

    def fixture(self, labels=None):
        book = Workbook()
        a = book.active
        a.title = '甲'
        b = book.create_sheet('乙')
        labels = labels or [['收入', '成本', '利润', '现金', '期末'], ['收入', '利润', '现金', '期末']]
        for ws, names in zip([a, b], labels):
            ws['B4'] = '项目名称'
            for i, label in enumerate(names, 5):
                ws.cell(i, 2, label)
                ws.cell(i, 3, i * 10)
        book.save(self.source)
        book.close()

    def test_clustering_and_plan(self):
        self.fixture()
        result = analyze(self.source)
        self.assertEqual(len(result['groups']), 1)
        group = result['groups'][0]
        self.assertFalse(group['blockers'])
        self.assertEqual([(a['sheet'], a['old_row'], a['label']) for a in group['actions']], [('乙', 6, '成本')])
        book = load_workbook(self.source)
        unrelated = book.create_sheet('无关')
        unrelated.append(['', '项目名称'])
        for label in ['苹果', '梨', '桃']:
            unrelated.append(['', label])
        book.save(self.source)
        book.close()
        self.assertEqual(len(analyze(self.source)['groups']), 2)

    def test_repeated_context_and_ambiguity(self):
        labels = ['分部一', '人员工资', '收入', '分部二', '人员工资', '利润']
        self.fixture([labels, labels])
        self.assertFalse(analyze(self.source)['groups'][0]['actions'])
        self.assertFalse(analyze(self.source)['groups'][0]['blockers'])
        labels = ['分部一', '人员工资', '人员工资', '收入', '利润']
        self.fixture([labels, labels])
        result = analyze(self.source)
        self.assertTrue(result['groups'][0]['blockers'])
        with self.assertRaisesRegex(ViewerError, '歧义'):
            apply(self.source, self.output, result, 0, [])
        self.assertFalse(self.output.exists())

    def test_constants_formulas_properties_and_source_hash(self):
        self.fixture()
        book = load_workbook(self.source)
        b = book['乙']
        b['D6'] = '=C6+$C$5'
        b['I6'] = '=SUM(甲:乙!C5:C5)'
        book['甲']['D5'] = "='乙'!C6+SUM('乙'!C5:C8)"
        b['E6'] = True
        b['F6'] = '数据'
        b['F6'].comment = Comment('批注', '作者')
        b['F6'].hyperlink = 'https://example.org'
        b['F6'].font = Font(bold=True)
        b.row_dimensions[6].height = 28
        b.row_dimensions[6].hidden = True
        b.merge_cells('G6:H6')
        b['G6'] = '合并'
        book.save(self.source)
        book.close()
        before = fingerprint(self.source)
        result = analyze(self.source)
        report = apply(self.source, self.output, result, 0, [0])
        self.assertTrue(report['passed'])
        self.assertEqual(before, fingerprint(self.source))
        out = load_workbook(self.output)
        self.addCleanup(out.close)
        self.assertEqual(out['乙']['D7'].value, '=C7+$C$5')
        self.assertEqual(out['乙']['I7'].value, '=SUM(甲:乙!C5:C5)')
        self.assertEqual(out['甲']['D5'].value, "='乙'!C7+SUM('乙'!C5:C9)")
        self.assertIs(out['乙']['E7'].value, True)
        self.assertEqual(out['乙']['F7'].comment.text, '批注')
        self.assertEqual(out['乙']['F7'].hyperlink.ref, 'F7')
        self.assertTrue(out['乙']['F7'].font.bold)
        self.assertEqual(out['乙'].row_dimensions[7].height, 28)
        self.assertTrue(out['乙'].row_dimensions[7].hidden)
        self.assertIn('G7:H7', str(out['乙'].merged_cells))
        self.assertIsNone(out['乙']['C6'].value)

    def test_save_as_and_stale_source(self):
        self.fixture()
        with self.assertRaisesRegex(ViewerError, '源文件'):
            validate_destination(self.source, self.source)
        self.output.write_bytes(b'existing')
        with self.assertRaisesRegex(ViewerError, '已存在'):
            validate_destination(self.source, self.output)
        self.output.unlink()
        result = analyze(self.source)
        book = load_workbook(self.source)
        book['甲']['Z1'] = 123
        book.save(self.source)
        book.close()
        with self.assertRaisesRegex(ViewerError, '已变化'):
            apply(self.source, self.output, result, 0, [0])
        self.assertFalse(self.output.exists())

    def test_partial_selection_and_consecutive_insertions(self):
        complete = ['收入', '成本', '费用', '利润', '现金', '期末', '资产', '负债', '权益', '总计']
        short = [label for label in complete if label not in {'成本', '费用'}]
        self.fixture([complete, short])
        result = analyze(self.source)
        self.assertEqual(len(result['groups'][0]['actions']), 2)
        report = apply(self.source, self.output, result, 0, [1])
        self.assertFalse(report['full_union'])
        book = load_workbook(self.output)
        self.assertEqual(book['乙']['B6'].value, '费用')
        self.assertEqual(book['乙']['B7'].value, '利润')
        book.close()
        self.output.unlink()
        apply(self.source, self.output, result, 0, [0, 1])
        book = load_workbook(self.output)
        self.assertEqual([book['乙'].cell(r, 2).value for r in range(5, 15)], complete)
        book.close()

    def test_verification_failure_removes_only_temporary_output(self):
        self.fixture()
        result = analyze(self.source)
        before = fingerprint(self.source)
        original_save = Workbook.save

        def damaged_save(book, path):
            book['乙']['C7'] = 99999
            original_save(book, path)

        with patch.object(Workbook, 'save', damaged_save):
            with self.assertRaisesRegex(ViewerError, '验证失败'):
                apply(self.source, self.output, result, 0, [0])
        self.assertFalse(self.output.exists())
        self.assertFalse(list(self.source.parent.glob('.alignment-*')))
        self.assertEqual(fingerprint(self.source), before)

    def test_unsupported_formula_blocks_without_output(self):
        self.fixture()
        book = load_workbook(self.source)
        book['乙']['D6'] = '=INDIRECT("C6")'
        book.save(self.source)
        book.close()
        result = analyze(self.source)
        with self.assertRaisesRegex(ViewerError, '公式无法安全'):
            apply(self.source, self.output, result, 0, [0])
        self.assertFalse(self.output.exists())
        self.assertEqual(list(self.source.parent.glob('.alignment-*')), [])

    def test_subset_is_trusted_and_persisted(self):
        self.fixture()
        book = load_workbook(self.source)
        ws = book.copy_worksheet(book['甲'])
        ws.title = '冲突'
        ws['B6'], ws['B7'] = '利润', '成本'
        book.save(self.source)
        book.close()
        self.assertTrue(analyze(self.source)['groups'][0]['blockers'])
        scope = {'0': {'sheets': ['甲', '乙']}}
        result = analyze(self.source, scope)
        self.assertFalse(result['groups'][0]['blockers'])
        self.assertEqual(len(result['groups'][0]['actions']), 1)
        result['groups'][0]['actions'][0]['old_row'] = 999
        with self.assertRaisesRegex(ViewerError, '过期'):
            apply(self.source, self.output, result, 0, [0])
        result = analyze(self.source, scope)
        apply(self.source, self.output, result, 0, [0])
        with self.assertRaisesRegex(ViewerError, '子集'):
            analyze(self.source, {'0': {'sheets': ['不存在', '乙']}})
        self.assertTrue(analyze(self.source, {'0': {'sheets': ['甲']}})['groups'][0]['blockers'])

    def test_explicit_equivalence_preserves_source_labels(self):
        base = ['收入', '成本', '费用', '现金', '期末', '资产', '负债', '权益', '总计']
        variant = [x for x in base if x != '成本']
        variant[variant.index('费用')] = '支出'
        self.fixture([base, variant])
        group = analyze(self.source)['groups'][0]
        pair = next(c['labels'] for c in group['candidates'] if set(c['labels']) == {'费用', '支出'})
        result = analyze(self.source, {'0': {'sheets': ['甲', '乙'], 'equivalences': [pair]}})
        self.assertFalse(result['groups'][0]['blockers'])
        apply(self.source, self.output, result, 0, [0])
        book = load_workbook(self.output)
        self.assertEqual(book['乙']['B6'].value, '成本')
        self.assertEqual(book['乙']['B7'].value, '支出')
        self.assertEqual(book['甲']['B7'].value, '费用')
        book.close()

    def test_unrelated_formulas_and_features_survive_save(self):
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.formatting.rule import CellIsRule
        self.fixture()
        book = load_workbook(self.source)
        ws = book.create_sheet('无关')
        formulas = ["=SUM('[1]外部'!A:A)", '=SUM(A:A)', "=SUM('乙'!C:C)+'乙'!C6",
                    '=INDIRECT("A1")', '=SUM(1:3)', '=SUM(甲:乙!C5:C5)']
        for i, formula in enumerate(formulas, 1):
            ws.cell(i, 4, formula)
        dv = DataValidation(type='whole', operator='greaterThan', formula1=0)
        ws.add_data_validation(dv)
        dv.add('A1:A10')
        ws.auto_filter.ref = 'A1:B10'
        ws.conditional_formatting.add('B1:B10', CellIsRule(operator='greaterThan', formula=['0']))
        book.save(self.source)
        book.close()
        result = analyze(self.source)
        self.assertFalse(result['groups'][0]['blockers'])
        apply(self.source, self.output, result, 0, [0])
        book = load_workbook(self.output)
        ws = book['无关']
        formulas[2] = "=SUM('乙'!C:C)+'乙'!C7"
        self.assertEqual([ws.cell(i, 4).value for i in range(1, 7)], formulas)
        self.assertEqual(str(ws.data_validations.dataValidation[0].sqref), 'A1:A10')
        self.assertEqual(ws.auto_filter.ref, 'A1:B10')
        self.assertEqual(str(next(iter(ws.conditional_formatting)).sqref), 'B1:B10')
        book.close()

    def test_affected_feature_and_unsupported_cross_reference_block(self):
        self.fixture()
        book = load_workbook(self.source)
        book['乙'].auto_filter.ref = 'B4:C10'
        book['甲']['D5'] = "=SUM('乙'!6:8)"
        book['甲']['E5'] = '=SUM(甲:乙!C6:C6)'
        book.save(self.source)
        book.close()
        blockers = analyze(self.source)['groups'][0]['blockers']
        self.assertTrue(any('筛选' in b for b in blockers))
        self.assertTrue(any('公式无法安全' in b for b in blockers))
        self.assertTrue(any('受影响的三维引用' in b for b in blockers))

    def test_terminology_and_order_conflicts_block(self):
        base = ['收入', '成本', '费用', '现金', '期末', '资产', '负债', '权益', '总计']
        variant = base.copy()
        variant[2] = '支出'
        self.fixture([base, variant])
        self.assertTrue(analyze(self.source)['groups'][0]['blockers'])
        variant = base.copy()
        variant[2:4] = reversed(variant[2:4])
        self.fixture([base, variant])
        self.assertTrue(analyze(self.source)['groups'][0]['blockers'])


if __name__ == '__main__':
    unittest.main()
