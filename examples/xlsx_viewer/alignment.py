"""Conservative insertion-only alignment and verified, atomic Save As worker."""
from collections import Counter
from copy import copy, deepcopy
from difflib import SequenceMatcher
import hashlib
import json
import os
from pathlib import Path
import re
import sys
import tempfile

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import get_column_letter
from openpyxl.worksheet.cell_range import CellRange

from loader import ViewerError, read_workbook, MAX_ROWS, _display_value

WARNING = "openpyxl 无法保证所有 Excel 对象完整往返；绘图、外部链接等可能变化。请保留原件并在 Excel 中复核。"


def fingerprint(path):
    with open(path, 'rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest() if hasattr(hashlib, 'file_digest') else hashlib.sha256(stream.read()).hexdigest()


def descriptors(sheet):
    rows = sheet['rows']
    headers = [(r, c) for r, row in enumerate(rows[:20])
               for c, value in enumerate(row[:10]) if value.strip() == '项目名称']
    if len(headers) != 1:
        return None
    header, col = headers[0]
    labels = [(r + 1, row[col].strip()) for r, row in enumerate(rows)
              if r > header and len(row) > col and row[col].strip()]
    counts = Counter(label for _, label in labels)
    anchor = '<表头>'
    result = []
    for row, label in labels:
        # Repeated labels are qualified by the nearest unique preceding label.
        key = (label, anchor if counts[label] > 1 else '')
        result.append({'row': row, 'label': label, 'key': list(key), 'context': anchor})
        if counts[label] == 1:
            anchor = label
    return {'name': sheet['name'], 'header': header + 1, 'column': col + 1, 'rows': result}


def similarity(a, b):
    x, y = [r['label'] for r in a['rows']], [r['label'] for r in b['rows']]
    return SequenceMatcher(None, x, y, autojunk=False).ratio() if min(len(x), len(y)) >= 3 else 0


def keyed_sheets(sheets, equivalences=None):
    sheets = deepcopy(sheets)
    canonical = {}
    for first, second in equivalences or []:
        a, b = canonical.get(first, first), canonical.get(second, second)
        canonical = {k: a if v == b else v for k, v in canonical.items()}
        canonical[first] = canonical[second] = a
    for sheet in sheets:
        counts = Counter(canonical.get(r['label'], r['label']) for r in sheet['rows'])
        anchor = '<表头>'
        for row in sheet['rows']:
            label = canonical.get(row['label'], row['label'])
            row['key'] = [label, anchor if counts[label] > 1 else '']
            if counts[label] == 1:
                anchor = label
    return sheets


def make_plan(sheets, equivalences=None):
    sheets = keyed_sheets(sheets, equivalences)
    candidates = []
    nodes, edges, order = {}, {}, []
    blockers = []
    for sheet in sheets:
        keys = [tuple(r['key']) for r in sheet['rows']]
        if len(keys) != len(set(keys)):
            blockers.append(f"{sheet['name']}：重复标签的上下文仍不唯一，需人工整理。")
        for row, key in zip(sheet['rows'], keys):
            if key not in nodes:
                nodes[key] = row
                order.append(key)
                edges[key] = set()
        for a, b in zip(keys, keys[1:]):
            edges[a].add(b)
    union = []
    remaining = set(order)
    while remaining:
        ready = [k for k in order if k in remaining and not any(k in edges[p] for p in remaining)]
        if not ready:
            blockers.append('标签顺序冲突：本版本不自动重排行，请人工确认顺序。')
            break
        # Unrelated alternatives at the same gap might be synonyms: do not guess.
        if len(ready) > 1:
            blockers.append('同一位置存在多个候选标签（可能为术语变体）：' + ' / '.join(k[0] for k in ready))
            for i, a in enumerate(ready):
                for b in ready[i + 1:]:
                    if a[0] != b[0]:
                        candidates.append({'labels': [a[0], b[0]], 'reason': '同一顺序空隙中的不同标签；仅由用户确认等价'})
        key = ready[0]
        union.append(key)
        remaining.remove(key)
    actions = []
    if len(union) == len(nodes):
        for sheet in sheets:
            by_key = {tuple(r['key']): r for r in sheet['rows']}
            added = 0
            for index, key in enumerate(union):
                if key in by_key:
                    continue
                following = next((by_key[k]['row'] for k in union[index+1:] if k in by_key), None)
                old_row = following if following is not None else (sheet['rows'][-1]['row'] + 1 if sheet['rows'] else sheet['header'] + 1)
                actions.append({'id': len(actions), 'sheet': sheet['name'], 'type': 'insert',
                                'old_row': old_row, 'new_row': old_row + added,
                                'column': sheet['column'], 'label': key[0],
                                'context': nodes[key]['context'], 'confidence': 1.0,
                                'reason': '精确标签及上下文的顺序并集；仅新增标签与空白模板行',
                                'selected': not blockers, 'blocking': bool(blockers)})
                added += 1
    return {'sheets': [s['name'] for s in sheets], 'actions': actions, 'blockers': blockers, 'candidates': candidates}


def evidence_index(book, sheets):
    """Original rows only. Never evaluate formulas or infer business semantics."""
    index = {}
    for sheet in sheets:
        ws = book[sheet['name']]
        for row in sheet['rows']:
            cells = []
            for cell in ws[row['row']][sheet['column']:]:
                if cell.value is None:
                    continue
                formula = cell.data_type == 'f'
                numeric = cell.data_type == 'n' and isinstance(cell.value, (int, float))
                cells.append({'cell': cell.coordinate, 'value': _display_value(cell.value),
                              'header': ' / '.join(_display_value(ws.cell(r, cell.column).value)
                                                  for r in range(1, sheet['header'] + 1)
                                                  if ws.cell(r, cell.column).value is not None),
                              'formula': formula,
                              'nonzero': cell.value != 0 if numeric else None})
            evidence = {'sheet': ws.title, 'row': row['row'], 'label': row['label'],
                        'context': row['context'],
                        'cell': f"{get_column_letter(sheet['column'])}{row['row']}",
                        'description': _display_value(ws.cell(row['row'], 1).value)
                                       if sheet['column'] != 1 else '',
                        'has_nonzero': any(c['nonzero'] is True for c in cells),
                        'has_formula': any(c['formula'] for c in cells), 'cells': cells}
            index.setdefault(row['label'], []).append(evidence)
    return index


def order_advisories(sheets, index):
    """Report actionable inversions; fall back to one explicit multi-sheet cycle."""
    positions = [{tuple(r['key']): r for r in s['rows']} for s in sheets]
    # Ambiguous duplicate keys cannot cast a reliable order vote.
    usable = [(s, p) for s, p in zip(sheets, positions) if len(p) == len(s['rows'])]
    pairs = set()
    for i, (_, first) in enumerate(usable):
        for _, second in usable[i + 1:]:
            common = [k for k in first if k in second]
            for a, b in zip(common, common[1:]):
                if second[a]['row'] > second[b]['row']:
                    pairs.add(tuple(sorted((a, b))))
    result = []
    for a, b in sorted(pairs):
        votes = []
        for sheet, pos in usable:
            if a in pos and b in pos:
                votes.append({'sheet': sheet['name'], 'rows': [pos[a]['row'], pos[b]['row']],
                              'forward': pos[a]['row'] < pos[b]['row']})
        forward = sum(v['forward'] for v in votes)
        reverse = len(votes) - forward
        majority = None if forward == reverse else [a, b] if forward > reverse else [b, a]
        moves = []
        if majority:
            for sheet, pos in usable:
                x, y = majority
                if x in pos and y in pos and pos[x]['row'] > pos[y]['row']:
                    moves.append(f"{sheet['name']}：建议人工核对后将原行 {pos[x]['row']}"
                                 f"（{pos[x]['label']}）移到原行 {pos[y]['row']}（{pos[y]['label']}）之前")
        evidence = [e for key in (a, b) for s, p in usable if key in p
                    for e in index.get(p[key]['label'], [])
                    if e['sheet'] == s['name'] and e['row'] == p[key]['row']]
        result.append({'kind': 'order', 'labels': [a[0], b[0]], 'keys': [list(a), list(b)],
                       'votes': votes, 'majority': [list(k) for k in majority] if majority else None,
                       'evidence': evidence,
                       'recommendation': ('；'.join(moves) if moves else '顺序票数相同，请人工确认') +
                       '。多数顺序仅供参考，不证明业务正确；非空行不得按缺项处理，不自动移动。'})
    if result:
        return result
    # A->B->C->A can exist without any pair sharing two sheets.
    edges = {}
    for sheet, _ in usable:
        keys = [tuple(r['key']) for r in sheet['rows']]
        for key in keys:
            edges.setdefault(key, [])
        for a, b in zip(keys, keys[1:]):
            if b not in edges[a]:
                edges[a].append(b)
    done = set()
    for root in edges:
        if root in done:
            continue
        path, active = [root], {root: 0}
        stack = [iter(edges[root])]
        while stack:
            node = next(stack[-1], None)
            if node is None:
                finished = path.pop()
                done.add(finished)
                del active[finished]
                stack.pop()
            elif node in active:
                cycle = path[active[node]:] + [node]
                links = []
                for a, b in zip(cycle, cycle[1:]):
                    links.append({'keys': [list(a), list(b)], 'sources': [
                        {'sheet': s['name'], 'rows': [p[a]['row'], p[b]['row']]}
                        for s, p in usable if a in p and b in p and p[a]['row'] < p[b]['row']]})
                return [{'kind': 'order_cycle', 'labels': [k[0] for k in cycle],
                         'links': links, 'evidence': [],
                         'recommendation': '跨表顺序形成循环，缺少可用的成对多数顺序；请按列出的原行人工确认，不自动移动。'}]
            elif node not in done:
                active[node] = len(path)
                path.append(node)
                stack.append(iter(edges[node]))
    return []


def difference_advisories(plan, chosen, index):
    result = []
    for action in plan['actions']:
        instances = index.get(action['label'], [])
        target = [e for e in instances if e['sheet'] == action['sheet']]
        recommendation = ('目标表存在同名行，可能是上下文不同；请人工核对，不能据此判定缺少业务数据。'
                          if target else
                          '源表未发现该主体业务数据（仅限目标表可识别项目名称列中的同名字段）；不能排除异名或未识别区域，只建议空白占位，禁止复制别家值。')
        result.append({'kind': 'insertion', 'action_id': action['id'],
                       'labels': [action['label']], 'target_sheet': action['sheet'],
                       'target_has_same_label': bool(target), 'target_instances': target,
                       'evidence': instances, 'recommendation': recommendation})
    for candidate in plan['candidates']:
        instances = [e for label in candidate['labels'] for e in index.get(label, [])]
        nonzero = any(e['has_nonzero'] for e in instances)
        result.append({'kind': 'terminology', 'labels': candidate['labels'],
                       'evidence': instances,
                       'recommendation': ('存在非零业务数据，标签可能代表语义不同的独立业务；默认按不等价处理，不建议视为同义词，需人工确认。'
                                          if nonzero else
                                          '标签不同；公式结果或业务语义未确认，默认按不等价处理，需人工确认。') +
                       '同一位置不构成等价证据；描述、业务分类有冲突时不得自动确认。'})
    result.extend(order_advisories(keyed_sheets(chosen, plan['equivalences']), index))
    return result


def analyze(path, scopes=None):
    before = fingerprint(path)
    candidates = [d for s in read_workbook(path) if (d := descriptors(s)) and d['rows']]
    groups = []
    for sheet in candidates:
        group = next((g for g in groups if all(similarity(sheet, member) >= (.9 if max(len(sheet['rows']), len(member['rows'])) > 10 else .8) for member in g)), None)
        if group is None:
            groups.append([sheet])
        else:
            group.append(sheet)
    if fingerprint(path) != before:
        raise ViewerError('分析期间源文件已变化，请重新打开。')
    plans = []
    scopes = scopes or {}
    # Analysis needs formulas and local features, not external cached cell data.
    book = load_workbook(path, keep_links=False)
    try:
        evidence = evidence_index(book, candidates)
        for index, members in enumerate(groups):
            scope = scopes.get(str(index), {})
            names = scope.get('sheets', [s['name'] for s in members])
            if len(names) != len(set(names)) or not set(names) <= {s['name'] for s in members}:
                raise ViewerError('无效工作表子集。')
            chosen = [s for s in members if s['name'] in names]
            equivalences = scope.get('equivalences', [])
            # Only approve pairs surfaced by the trusted original subset plan.
            available = [c['labels'] for c in make_plan(chosen)['candidates']]
            if any(pair not in available for pair in equivalences):
                raise ViewerError('无效术语等价候选。')
            plan = make_plan(chosen, equivalences)
            plan.update(members=[s['name'] for s in members], equivalences=equivalences)
            plan['candidates'] = make_plan(chosen)['candidates']
            if len(chosen) < 2:
                plan['blockers'].append('请至少选择两张候选工作表。')
            if plan['actions'] and not plan['blockers']:
                plan['blockers'].extend(preflight(book, plan['actions']))
            for action in plan['actions']:
                action.update(selected=not plan['blockers'], blocking=bool(plan['blockers']))
            plan['advisories'] = difference_advisories(plan, chosen, evidence)
            plans.append(plan)
    finally:
        book.close()
    if fingerprint(path) != before:
        raise ViewerError('分析期间源文件已变化。')
    return {'sha256': before, 'groups': plans, 'scopes': scopes, 'warning': WARNING}


CELL = r'\$?[A-Za-z]{1,3}\$?[1-9][0-9]*'
REFERENCE = re.compile(rf'(?:(\x27(?:[^\x27]|\x27\x27)+\x27|[^!]+)!)?({CELL})(?::({CELL}))?\Z')


def rewrite_formula(value, sheet, origin, target, maps, names):
    moving = origin != target
    affected = {name.casefold(): items for name, items in maps.items() if items}
    if not isinstance(value, str):
        raise ViewerError(f'{sheet}!{origin}：不支持数组/数据表公式。')
    try:
        tokens = Tokenizer(value).items
        unsupported = []
        changed = False
        dynamic = False
        for token in tokens:
            if token.type == 'FUNC' and token.value.upper().rsplit('.', 1)[-1] in {'INDIRECT(', 'OFFSET('}:
                dynamic = True
            if token.type != 'OPERAND' or token.subtype != 'RANGE':
                continue
            qualifier, sep, address = token.value.rpartition('!')
            referenced = (qualifier.strip("'").replace("''", "'") if sep else sheet)
            address = address if sep else token.value
            if '[' in referenced:  # External workbook references never follow local rows.
                if moving:
                    unsupported.append('移动单元格中的外部引用')
                continue
            if ':' in referenced:
                ends = referenced.split(':')
                lower_names = [n.casefold() for n in names]
                region = re.fullmatch(rf'({CELL})(?::({CELL}))?', address)
                if len(ends) != 2 or any(n.casefold() not in lower_names for n in ends) or not region:
                    unsupported.append('无法判定的三维引用')
                    continue
                start, end = [lower_names.index(n.casefold()) for n in ends]
                rows = [int(re.search(r'[0-9]+$', a).group()) for a in region.groups() if a]
                if any(mapped_row(affected.get(n, []), r) != r
                                 for n in lower_names[min(start, end):max(start, end) + 1] for r in rows):
                    unsupported.append('受影响的三维引用')
                continue
            entries = affected.get(referenced.casefold(), [])
            match = REFERENCE.fullmatch(token.value)
            if match and referenced.casefold() in {n.casefold() for n in names}:
                _, first, last = match.groups()
                def remap(addr):
                    col, absolute, row = re.fullmatch(r'(\$?[A-Za-z]+)(\$?)([0-9]+)', addr).groups()
                    return col + absolute + str(mapped_row(entries, int(row)))
                replacement = (qualifier + '!' if sep else '') + remap(first) + (':' + remap(last) if last else '')
                changed |= replacement != token.value
                token.value = replacement
            elif re.fullmatch(r'\$?[A-Za-z]{1,3}:\$?[A-Za-z]{1,3}', address):
                # Whole columns have no row boundary to adjust.
                if moving:
                    unsupported.append('移动单元格中的整列引用')
            elif re.fullmatch(r'\$?[1-9][0-9]*:\$?[1-9][0-9]*', address):
                rows = [int(x.replace('$', '')) for x in address.split(':')]
                if moving or any(mapped_row(entries, r) != r for r in rows):
                    unsupported.append('受影响的整行引用')
            elif entries or moving or not sep or ':' in referenced:
                unsupported.append('名称、三维或非 A1 引用')
        if dynamic:
            # OFFSET follows its explicit A1 operand. INDIRECT literals can be
            # scoped; computed text remains conservatively unresolved.
            indirect = re.findall(r'INDIRECT\(\s*"([^"]+)"\s*\)', value, re.I)
            indirect_count = len(re.findall(r'INDIRECT\(', value, re.I))
            if moving or changed or indirect_count != len(indirect) or (re.search(r'OFFSET\(', value, re.I) and affected):
                unsupported.append('动态地址函数')
            for address in indirect:
                probe = rewrite_formula('=' + address, sheet, origin, origin, maps, names)
                if probe != '=' + address:
                    unsupported.append('动态地址函数指向移动行')
        if unsupported:
            raise ValueError(' / '.join(sorted(set(unsupported))))
        return '=' + ''.join(t.value for t in tokens) if changed else value
    except Exception as exc:
        raise ViewerError(f'{sheet}!{origin}：公式无法安全调整（{exc}）。') from exc


def preflight(book, actions):
    maps = {a['sheet']: [x for x in actions if x['sheet'] == a['sheet']] for a in actions}
    blockers = []
    if book.defined_names:
        blockers.append('工作簿包含定义名称，当前版本不能安全调整。')
    for ws in book.worksheets:
        entries = maps.get(ws.title, [])
        if entries and (ws.defined_names or ws.tables or ws.data_validations.count or
                        len(ws.conditional_formatting) or ws.auto_filter.ref or ws.print_area or
                        ws.print_title_rows or ws.print_title_cols):
            blockers.append(f'{ws.title}：受影响表含名称/表格/验证/条件格式/筛选/打印区域。')
        for row in ws:
            for cell in row:
                if cell.data_type == 'f':
                    target = f'{get_column_letter(cell.column)}{mapped_row(entries, cell.row)}'
                    try:
                        rewrite_formula(cell.value, ws.title, cell.coordinate, target, maps, book.sheetnames)
                    except ViewerError as exc:
                        blockers.append(str(exc))
        for merged in ws.merged_cells.ranges:
            if any(merged.min_row < a['old_row'] <= merged.max_row for a in entries):
                blockers.append(f'{ws.title}!{merged}：插入位置穿过合并区域。')
    return blockers


def mapped_row(insertions, row):
    return row + sum(a['old_row'] <= row for a in insertions)


def validate_destination(source, destination):
    source, destination = Path(source).resolve(), Path(destination).resolve()
    immutable = Path(__file__).resolve().parents[2] / 'data'
    if destination == source or (destination.exists() and os.path.samefile(source, destination)):
        raise ViewerError('另存路径不能是源文件。')
    if destination.is_relative_to(immutable):
        raise ViewerError('data/ 是只读输入目录，请选择其他输出目录。')
    if destination.suffix.lower() != source.suffix.lower():
        raise ViewerError('输出扩展名必须与源文件一致（保留 VBA）。')
    if destination.exists():
        raise ViewerError('输出已存在，请选择新的文件名。')
    return destination


def apply(source, destination, analysis, group_index, selected, staging=None):
    destination = validate_destination(source, destination)
    if fingerprint(source) != analysis['sha256']:
        raise ViewerError('源文件已变化，请重新分析。')
    # Recompute trusted actions; never accept client-supplied row coordinates.
    current = analyze(source, analysis.get('scopes'))
    if current != analysis:
        raise ViewerError('分析结果已过期。')
    group = current['groups'][group_index]
    if group['blockers']:
        raise ViewerError('存在阻断歧义：' + '；'.join(group['blockers']))
    actions = [a for a in group['actions'] if a['id'] in selected]
    if not actions or len(actions) != len(set(selected)):
        raise ViewerError('请选择有效操作。')
    maps = {name: [a for a in actions if a['sheet'] == name] for name in group['sheets']}
    book = load_workbook(source, keep_vba=Path(source).suffix.lower() == '.xlsm', keep_links=True)
    temp = None
    try:
        blockers = preflight(book, actions)
        if blockers:
            raise ViewerError('；'.join(blockers))
        expected = {}
        constants = formulas = 0
        for ws in book.worksheets:
            insertions = maps.get(ws.title, [])
            if ws.max_row + len(insertions) > MAX_ROWS:
                raise ViewerError('插入后超过行数限制。')
            for row in ws:
                for cell in row:
                    if cell.value is None:
                        continue
                    dest = f'{get_column_letter(cell.column)}{mapped_row(insertions, cell.row)}'
                    value = cell.value
                    if cell.data_type == 'f':
                        value = rewrite_formula(value, ws.title, cell.coordinate, dest, maps, book.sheetnames)
                        formulas += 1
                    else:
                        constants += 1
                    expected[(ws.title, dest)] = (value, type(value), cell.data_type)
            merges = [str(m) for m in ws.merged_cells.ranges]
            dimensions = [(r, copy(d)) for r, d in ws.row_dimensions.items()]
            for merged in list(ws.merged_cells.ranges):
                if any(merged.min_row < a['old_row'] <= merged.max_row for a in insertions):
                    raise ViewerError(f'{ws.title}!{merged}：插入位置穿过合并区域。')
            if not insertions:
                continue
            for merged in merges:
                ws.unmerge_cells(merged)
            # Snapshot styles for inserted empty template cells before movement.
            templates = {a['id']: [copy(c._style) for c in ws[a['old_row']]] for a in insertions}
            for a in sorted(insertions, key=lambda a: (a['old_row'], a['id']), reverse=True):
                ws.insert_rows(a['old_row'])
            ws.row_dimensions.clear()
            for r, dimension in dimensions:
                new = mapped_row(insertions, r)
                dimension.index = new
                ws.row_dimensions[new] = dimension
            for merged in merges:
                area = CellRange(merged)
                ws.merge_cells(start_row=mapped_row(insertions, area.min_row),
                               end_row=mapped_row(insertions, area.max_row),
                               start_column=area.min_col, end_column=area.max_col)
            for offset, a in enumerate(sorted(insertions, key=lambda a: (a['old_row'], a['id']))):
                r = a['old_row'] + offset
                for c, style in enumerate(templates[a['id']], 1):
                    ws.cell(r, c)._style = style
                cell = ws.cell(r, a['column'])
                cell.value = a['label']
                cell.data_type = 's'
                expected[(ws.title, cell.coordinate)] = (cell.value, type(cell.value), cell.data_type)
            # insert_rows moves Cell objects but hyperlink refs retain old addresses.
            for row in ws:
                for cell in row:
                    if cell.hyperlink:
                        cell.hyperlink.ref = cell.coordinate
        for (name, coordinate), (value, _, kind) in expected.items():
            if kind == 'f':
                book[name][coordinate] = value
        if staging:
            candidate = Path(staging)
            if (candidate.is_symlink() or candidate.resolve().parent != destination.parent
                    or not candidate.name.startswith('.alignment-')
                    or candidate.suffix != destination.suffix
                    or not candidate.is_file() or candidate.stat().st_size != 0):
                raise ViewerError('无效暂存路径。')
            temp = candidate
        else:
            fd, name = tempfile.mkstemp(prefix='.alignment-', suffix=destination.suffix, dir=destination.parent)
            os.close(fd)
            temp = Path(name)
        book.save(temp)
        check = load_workbook(temp, keep_vba=destination.suffix.lower() == '.xlsm', keep_links=True)
        try:
            if check.sheetnames != book.sheetnames:
                raise ViewerError('验证失败：工作表集合改变。')
            actual = {(ws.title, c.coordinate): (c.value, type(c.value), c.data_type)
                      for ws in check.worksheets for row in ws for c in row if c.value is not None}
            if actual != expected:
                bad = [f'{s}!{c}' for s, c in expected.keys() | actual.keys()
                       if actual.get((s, c)) != expected.get((s, c))]
                raise ViewerError('验证失败：值、类型或公式不一致：' + ', '.join(bad[:15]))
            if fingerprint(source) != analysis['sha256']:
                raise ViewerError('验证失败：源文件在处理期间改变。')
        finally:
            check.close()
        report = {'passed': True, 'constants': constants, 'formulas': formulas,
                  'actions': len(actions), 'full_union': len(actions) == len(group['actions']),
                  'checks': ['工作表集合', '常量值及 Python 类型与唯一目标', '全部标签行', '公式文本与目标位置', '源文件 SHA256'],
                  'warning': WARNING}
        if not staging:
            validate_destination(source, destination)
            os.replace(temp, destination)
        temp = None
        return report
    finally:
        book.close()
        if temp is not None:
            temp.unlink(missing_ok=True)


def main():
    try:
        request = json.load(sys.stdin)
        if request['mode'] == 'analyze':
            result = analyze(request['source'], request.get('scopes'))
        else:
            result = apply(request['source'], request['destination'], request['analysis'],
                           request['group'], request['selected'], request['staging'])
        print(json.dumps({'result': result}, ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({'error': str(exc)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
