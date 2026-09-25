"""Run with unittest discovery; Qt uses the offscreen platform."""

import hashlib
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import time
import unittest
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

os.environ["QT_QPA_PLATFORM"] = "offscreen"

from openpyxl import Workbook
from PySide6.QtCore import QProcess, Qt, QTimer
from PySide6.QtWidgets import QApplication, QFileDialog

import loader
from viewer import SheetModel, Viewer, column_title


class WorkbookFixture:
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "测试.xlsx"

    def make_book(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "数据"
        sheet.append(["你好", 42, True, "=B1*2"])
        sheet["C3"] = "end"
        book.create_sheet("空白")
        hidden = book.create_sheet("隐藏")
        hidden.sheet_state = "hidden"
        hidden["A1"] = "visible here"
        sheet.protection.sheet = True  # Editing protection is not encryption.
        book.save(self.path)
        book.close()


class WorkbookTests(WorkbookFixture, unittest.TestCase):
    def test_values_formulas_sheets_and_no_writes(self):
        self.make_book()
        before = hashlib.sha256(self.path.read_bytes()).digest()
        sheets = loader.read_workbook(self.path)
        self.assertEqual([s["name"] for s in sheets], ["数据", "空白", "隐藏"])
        self.assertEqual(sheets[0]["rows"][0], ["你好", "42", "True", "=B1*2"])
        self.assertEqual(sheets[0]["columns"], 4)
        self.assertEqual(sheets[0]["rows"][2][2], "end")
        self.assertEqual(sheets[1]["rows"], [])
        self.assertEqual(hashlib.sha256(self.path.read_bytes()).digest(), before)

    def test_xlsm_and_worker_protocol(self):
        self.make_book()
        # Macro-free OOXML fixture with the supported .xlsm extension.
        target = self.path.with_suffix(".xlsm")
        self.path.rename(target)
        result = subprocess.run(
            [sys.executable, str(Path(loader.__file__)), str(target)],
            capture_output=True, check=True, timeout=15,
        )
        self.assertEqual(json.loads(result.stdout)["sheets"][0]["name"], "数据")

    def test_invalid_files(self):
        for content, message in [(b"", "文件为空"), (b"not excel", "无法解析"),
                                 (loader.OLE_SIGNATURE + b"data", "密码加密")]:
            with self.subTest(content=content):
                self.path.write_bytes(content)
                with self.assertRaisesRegex(loader.ViewerError, message):
                    loader.read_workbook(self.path)
        with self.assertRaisesRegex(loader.ViewerError, "无法读取"):
            loader.read_workbook(self.path.with_name("missing.xlsx"))
        with self.assertRaisesRegex(loader.ViewerError, "仅支持"):
            loader.read_workbook(self.path.with_suffix(".xls"))
        with ZipFile(self.path, "w") as archive:
            archive.writestr("not-a-workbook.xml", "<root/>")
        with self.assertRaisesRegex(loader.ViewerError, "无法解析"):
            loader.read_workbook(self.path)

    def test_limits(self):
        self.make_book()
        for name, value, message in [
            ("MAX_FILE_BYTES", 1, "文件超过"),
            ("MAX_UNPACKED_BYTES", 1, "解压体积"),
            ("MAX_ZIP_ENTRIES", 1, "内部文件"),
            ("MAX_SHEETS", 1, "工作表超过"),
            ("MAX_ROWS", 2, "行或"),
            ("MAX_COLUMNS", 2, "列限制"),
            ("MAX_TEXT_CHARS", 2, "显示文本"),
        ]:
            with self.subTest(limit=name), patch.object(loader, name, value):
                with self.assertRaisesRegex(loader.ViewerError, message):
                    loader.read_workbook(self.path)

    def test_more_than_100000_cells_can_be_viewed(self):
        book = Workbook()
        sheet = book.active
        sheet["A1"] = "start"
        sheet.cell(row=400, column=256, value="end")
        book.save(self.path)
        book.close()

        sheets = loader.read_workbook(self.path)

        self.assertEqual(len(sheets[0]["rows"]) * sheets[0]["columns"], 102400)
        self.assertEqual(sheets[0]["rows"][399][255], "end")

    def test_understated_dimensions_do_not_bypass_limits(self):
        self.make_book()
        with ZipFile(self.path) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        name = "xl/worksheets/sheet1.xml"
        entries[name] = re.sub(rb'<dimension ref="[^"]+"',
                               b'<dimension ref="A1"', entries[name])
        with ZipFile(self.path, "w", ZIP_DEFLATED) as archive:
            for name, data in entries.items():
                archive.writestr(name, data)
        with patch.object(loader, "MAX_ROWS", 2):
            with self.assertRaisesRegex(loader.ViewerError, "行或"):
                loader.read_workbook(self.path)


class QtTests(WorkbookFixture, unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        super().setUp()
        self.window = Viewer()
        self.addCleanup(self.window.close)
        self.errors = []
        self.window._report_error = self.errors.append

    def wait_until(self, predicate):
        deadline = time.monotonic() + 10
        while not predicate() and time.monotonic() < deadline:
            self.app.processEvents()
            time.sleep(0.01)
        self.app.processEvents()
        self.assertTrue(predicate(), "Qt worker did not finish")

    def test_model_read_only_and_headers(self):
        model = SheetModel({"rows": [["=1+1"], ["a", "b"]], "columns": 2})
        self.assertEqual(model.rowCount(), 2)
        self.assertEqual(model.columnCount(), 2)
        self.assertEqual(model.data(model.index(0, 0)), "=1+1")
        self.assertEqual(model.data(model.index(0, 1)), "")
        self.assertFalse(model.flags(model.index(0, 0)) & Qt.ItemFlag.ItemIsEditable)
        self.assertFalse(model.setData(model.index(0, 0), "changed"))
        self.assertEqual(model.headerData(1, Qt.Orientation.Vertical), "2")
        self.assertEqual(model.headerData(26, Qt.Orientation.Horizontal), "AA")
        self.assertEqual([column_title(i) for i in [0, 25, 26, 255]],
                         ["A", "Z", "AA", "IV"])

    def test_open_switch_and_recover_from_error(self):
        self.make_book()
        with patch.object(QFileDialog, "getOpenFileName", return_value=(str(self.path), "")):
            self.window.choose_file()
        self.wait_until(lambda: self.window.open_button.isEnabled())
        self.assertEqual(self.errors, [])
        self.assertEqual(self.window.selector.count(), 3)
        self.assertEqual(self.window.model.data(self.window.model.index(0, 3)), "=B1*2")
        self.window.selector.setCurrentIndex(1)
        self.assertEqual(self.window.model.rowCount(), 0)
        self.assertIn("空工作表", self.window.status.text())
        self.window.selector.setCurrentIndex(2)
        self.assertEqual(self.window.model.data(self.window.model.index(0, 0)), "visible here")
        self.path.write_bytes(b"")
        self.window.load_file(self.path)
        self.wait_until(lambda: self.window.open_button.isEnabled())
        self.assertIn("文件为空", self.errors[-1])
        self.assertIsNone(self.window.table.model())
        self.make_book()
        self.window.load_file(self.path)
        self.wait_until(lambda: self.window.open_button.isEnabled())
        self.assertEqual(self.window.selector.count(), 3)

    def test_cancel_and_timeout_stop_worker(self):
        # Substitute a slow Python child so the timer test is deterministic.
        for action, expected in [(self.window.cancel_load, "已取消"),
                                 (self.window._timeout, "超过 30 秒")]:
            with self.subTest(action=action.__name__):
                self.window.failure = None
                self.window.open_button.setEnabled(False)
                self.window.process.start(sys.executable, ["-c", "import time; time.sleep(60)"])
                self.assertTrue(self.window.process.waitForStarted(3000))
                if action == self.window._timeout:
                    self.window.timer.start(20)
                else:
                    QTimer.singleShot(20, action)
                self.wait_until(lambda: self.window.open_button.isEnabled())
                self.assertEqual(self.window.process.state(), QProcess.ProcessState.NotRunning)
                self.assertIn(expected, self.window.status.text())

    def test_alignment_review_save_and_verified_output(self):
        book = Workbook()
        for i, labels in enumerate((["收入", "成本", "利润", "现金", "期末"],
                                    ["收入", "利润", "现金", "期末"])):
            sheet = book.active if i == 0 else book.create_sheet("乙")
            sheet["B4"] = "项目名称"
            for row, label in enumerate(labels, 5):
                sheet.cell(row, 2, label)
                sheet.cell(row, 3, row * 10)
        book.save(self.path)
        book.close()
        self.window.load_file(self.path)
        self.wait_until(lambda: self.window.open_button.isEnabled())
        self.assertFalse(self.window.save_button.isEnabled())
        self.window.start_analysis()
        self.wait_until(lambda: self.window.job is None)
        self.assertEqual(self.window.actions.rowCount(), 1)
        self.assertTrue(self.window.save_button.isEnabled())
        item = self.window.actions.item(0, 0)
        item.setCheckState(Qt.CheckState.Unchecked)
        self.assertFalse(self.window.save_button.isEnabled())
        item.setCheckState(Qt.CheckState.Checked)
        self.window.navigate_action(0, 0)
        self.assertEqual(self.window.selector.currentText(), "乙")
        destination = self.path.with_name("输出.xlsx")
        with patch.object(QFileDialog, "getSaveFileName", return_value=(str(destination), "")):
            self.window.save_aligned()
        self.wait_until(lambda: self.window.job is None)
        self.assertTrue(destination.exists(), self.window.summary.text())
        self.assertIn("已保存并验证", self.window.summary.text())
        self.assertIsNone(self.window.staging)

    def test_subset_regeneration_and_identical_state(self):
        book = Workbook()
        base = ["收入", "成本", "利润", "现金", "期末"]
        for i, labels in enumerate((base, [x for x in base if x != "成本"],
                                    ["收入", "利润", "成本", "现金", "期末"])):
            ws = book.active if i == 0 else book.create_sheet(str(i))
            ws["B4"] = "项目名称"
            for row, label in enumerate(labels, 5):
                ws.cell(row, 2, label)
        book.save(self.path)
        book.close()
        self.window.load_file(self.path)
        self.wait_until(lambda: self.window.open_button.isEnabled())
        self.window.start_analysis()
        self.wait_until(lambda: self.window.job is None)
        self.assertFalse(self.window.save_button.isEnabled())
        self.window.members.item(2, 0).setCheckState(Qt.CheckState.Unchecked)
        self.assertTrue(self.window.regenerate_button.isEnabled())
        self.window.regenerate_plan()
        self.wait_until(lambda: self.window.job is None)
        self.assertEqual(self.window.analysis['groups'][0]['sheets'], ['Sheet', '1'])
        self.assertEqual(self.window.actions.rowCount(), 1)
        self.assertTrue(self.window.save_button.isEnabled())
        self.window.members.item(1, 0).setCheckState(Qt.CheckState.Unchecked)
        self.assertFalse(self.window.regenerate_button.isEnabled())
        self.assertFalse(self.window.save_button.isEnabled())

    def test_identical_group_disables_save(self):
        book = Workbook()
        for ws in (book.active, book.create_sheet('相同')):
            ws.append(['项目名称'])
            for label in ['收入', '成本', '利润']:
                ws.append([label])
        book.save(self.path)
        book.close()
        self.window.load_file(self.path)
        self.wait_until(lambda: self.window.open_button.isEnabled())
        self.window.start_analysis()
        self.wait_until(lambda: self.window.job is None)
        self.assertIn('结构一致', self.window.summary.text())
        self.assertFalse(self.window.save_button.isEnabled())

    def test_alignment_cancel_cleans_staging(self):
        self.window.job = "apply"
        self.window.job_failure = None
        self.window.staging = self.path.with_name(".alignment-test.xlsx")
        self.window.staging.touch()
        staging = self.window.staging
        self.window.worker.start(sys.executable, ["-c", "import time; time.sleep(60)"])
        self.assertTrue(self.window.worker.waitForStarted(3000))
        self.window.cancel_job("已取消")
        self.wait_until(lambda: self.window.job is None)
        self.assertFalse(staging.exists())
        self.assertIn("已取消", self.window.summary.text())

    def test_close_stops_worker_without_error_dialog(self):
        self.window.process.start(sys.executable, ["-c", "import time; time.sleep(60)"])
        self.assertTrue(self.window.process.waitForStarted(3000))
        self.window.close()
        self.assertEqual(self.window.process.state(), QProcess.ProcessState.NotRunning)
        self.assertEqual(self.errors, [])


if __name__ == "__main__":
    unittest.main()
