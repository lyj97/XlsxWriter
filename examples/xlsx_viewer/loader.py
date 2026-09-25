"""Bounded, read-only workbook loading, also usable as a worker process."""

import json
from pathlib import Path
import sys
from zipfile import ZipFile

from openpyxl import load_workbook


MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_UNPACKED_BYTES = 100 * 1024 * 1024
MAX_ZIP_ENTRIES = 2000
MAX_SHEETS = 100
MAX_ROWS = 10000
MAX_COLUMNS = 256
MAX_TEXT_CHARS = 1000000
OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


class ViewerError(Exception):
    """An error suitable for display to the user."""


def _check_archive(source):
    source.seek(0, 2)
    size = source.tell()
    source.seek(0)
    if not size:
        raise ViewerError("文件为空（0 字节），请选择有效的 Excel 工作簿。")
    if size > MAX_FILE_BYTES:
        raise ViewerError("文件超过 20 MiB 限制，请先缩小工作簿。")
    if source.read(8) == OLE_SIGNATURE:
        raise ViewerError("不支持密码加密的工作簿或旧版 .xls；请用 Excel 另存为未加密的 .xlsx。")
    source.seek(0)
    with ZipFile(source) as archive:
        entries = archive.infolist()
        if any(entry.flag_bits & 1 for entry in entries):
            raise ViewerError("不支持密码加密的工作簿，请先另存为未加密的 .xlsx。")
        if (len(entries) > MAX_ZIP_ENTRIES or
                sum(entry.file_size for entry in entries) > MAX_UNPACKED_BYTES):
            raise ViewerError("工作簿解压体积超过 100 MiB 或内部文件超过 2000 个，无法查看。")
    source.seek(0)


def _display_value(value):
    if value is None:
        return ""
    # openpyxl represents array formulas as objects, with text on the object.
    if hasattr(value, "text"):
        return value.text or "[无可显示文本的公式]"
    if hasattr(value, "ref"):
        return "[数据表公式]"
    return str(value)


def read_workbook(filename):
    """Return sheet names and strings; never write or evaluate formulas."""
    path = Path(filename)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ViewerError("仅支持 .xlsx 和 .xlsm 文件。")
    try:
        # Use the same handle for inspection and loading, including on Windows.
        with path.open("rb") as source:
            _check_archive(source)
            workbook = load_workbook(
                source, read_only=True, data_only=False, keep_links=False
            )
            try:
                if not workbook.worksheets:
                    raise ViewerError("工作簿没有可显示的单元格工作表。")
                if len(workbook.worksheets) > MAX_SHEETS:
                    raise ViewerError("工作表超过 100 张限制。")
                sheets = []
                total_text = 0
                for sheet in workbook.worksheets:
                    if ((sheet.max_row or 0) > MAX_ROWS or
                            (sheet.max_column or 0) > MAX_COLUMNS):
                        raise ViewerError(f"工作表“{sheet.title}”超过 10000 行或 256 列限制。")
                    # Do not trust a producer's understated dimension metadata.
                    sheet.reset_dimensions()
                    rows = []
                    width = 0
                    for row in sheet.iter_rows():
                        width = max(width, len(row))
                        if len(rows) >= MAX_ROWS or width > MAX_COLUMNS:
                            raise ViewerError(f"工作表“{sheet.title}”超过 10000 行或 256 列限制。")
                        values = [_display_value(cell.value) for cell in row]
                        total_text += sum(len(value) for value in values)
                        if total_text > MAX_TEXT_CHARS:
                            raise ViewerError("工作簿显示文本超过 1000000 字符限制。")
                        rows.append(values)
                    sheets.append({"name": sheet.title, "rows": rows, "columns": width})
                return sheets
            finally:
                workbook.close()
    except ViewerError:
        raise
    except OSError as exc:
        raise ViewerError(f"无法读取文件：{exc}") from exc
    except Exception as exc:
        raise ViewerError("无法解析工作簿：文件可能损坏、被加密或不是有效的 .xlsx/.xlsm。") from exc


def main():
    try:
        result = {"sheets": read_workbook(sys.argv[1])}
    except ViewerError as exc:
        result = {"error": str(exc)}
    sys.stdout.buffer.write(json.dumps(result, ensure_ascii=False).encode("utf-8"))


if __name__ == "__main__":
    main()
